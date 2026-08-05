import asyncio
import logging
import time
from datetime import datetime, timezone
from io import BytesIO

import httpx
import openpyxl

from backend import storage
from backend.cache import registry

log = logging.getLogger("osint-globe.hdx_conflict_stats")

# HDX's own aggregate ACLED export (see acleddata.com's HDX org page) --
# country-month event/fatality counts, no registration or embargo, updated
# weekly. Complements backend/sources/acled.py's point-level feed, which on
# a Research-tier myACLED account is embargoed to events 12+ months old:
# this is the only near-current conflict signal available without a
# higher-tier account.
XLSX_URL = (
    "https://data.humdata.org/dataset/3e6bfc98-f837-495d-b8de-71e5ac026f59/"
    "resource/99a32d01-d0ca-4f57-a0f5-cb6b5f01f14f/download/"
    "political-violence-events-and-fatalities.xlsx"
)
POLL_INTERVAL = 6 * 3600  # the file itself only changes weekly; this just needs to notice within a day or so
FAILURE_RETRY_INTERVAL = 60  # scaled by consecutive failures, capped at POLL_INTERVAL
MONTHS_KEPT = 24  # bounds the payload -- this is a trend indicator, not an archive (full history lives in the source file itself)

_MONTH_NUM = {
    "January": 1, "February": 2, "March": 3, "April": 4, "May": 5, "June": 6,
    "July": 7, "August": 8, "September": 9, "October": 10, "November": 11, "December": 12,
}

# Country-level rows (Non_HRP) and admin2-level rows for humanitarian-crisis
# countries (HRP_1/HRP_2, summed back up to country totals) together cover
# every country ACLED tracks -- see the sheet layout noted in the getting-
# started research for this dataset.
_SHEETS = ("Non_HRP", "HRP_1", "HRP_2")


def _parse_workbook(raw: bytes) -> dict[str, list[dict]]:
    wb = openpyxl.load_workbook(BytesIO(raw), read_only=True)
    cutoff_ym = _months_ago(MONTHS_KEPT)

    # (country, year, month) -> {events, fatalities}, summed across admin
    # subdivisions for the HRP sheets so every country ends up with one
    # merged monthly series regardless of which sheet it came from.
    totals: dict[tuple[str, int, int], dict[str, int]] = {}

    for sheet_name in _SHEETS:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            continue
        idx = {name: i for i, name in enumerate(header)}
        try:
            ci, mi, yi, ei, fi = idx["Country"], idx["Month"], idx["Year"], idx["Events"], idx["Fatalities"]
        except KeyError:
            log.warning("HDX sheet %r missing an expected column, skipping", sheet_name)
            continue
        for row in rows:
            country, month_name, year = row[ci], row[mi], row[yi]
            month_num = _MONTH_NUM.get(month_name)
            if not country or not month_num or not year:
                continue
            year = int(year)
            if (year, month_num) < cutoff_ym:
                continue
            key = (country, year, month_num)
            bucket = totals.setdefault(key, {"events": 0, "fatalities": 0})
            bucket["events"] += int(row[ei] or 0)
            bucket["fatalities"] += int(row[fi] or 0)

    by_country: dict[str, list[dict]] = {}
    for (country, year, month_num), counts in totals.items():
        by_country.setdefault(country, []).append(
            {"year": year, "month": month_num, "events": counts["events"], "fatalities": counts["fatalities"]}
        )
    for series in by_country.values():
        series.sort(key=lambda r: (r["year"], r["month"]))
    return by_country


def _months_ago(n: int) -> tuple[int, int]:
    now = datetime.now(timezone.utc)
    year, month = now.year, now.month - n
    while month <= 0:
        month += 12
        year -= 1
    return (year, month)


async def _fetch() -> dict[str, list[dict]]:
    async with httpx.AsyncClient(timeout=120, follow_redirects=True) as client:
        resp = await client.get(XLSX_URL)
        resp.raise_for_status()
        raw = resp.content
    # openpyxl parsing this (a 40MB+ workbook, ~1M rows across sheets) is
    # blocking CPU work -- run it off the event loop so it doesn't stall
    # every other source's polling for the several seconds it takes.
    return await asyncio.to_thread(_parse_workbook, raw)


async def start():
    state = registry.register("hdx_conflict_stats", key_configured=True)
    consecutive_failures = 0
    while True:
        ok = False
        try:
            state.data = await _fetch()
            state.last_success = time.time()
            state.last_error = None
            ok = True
            log.info("HDX conflict stats: %d countries, last %d months", len(state.data), MONTHS_KEPT)
            # country -> monthly series dict, not point rows -- stored whole.
            await storage.record_reference("hdx_conflict_stats", state.data)
            await storage.record_source_health("hdx_conflict_stats", len(state.data), True)
        except Exception as exc:  # noqa: BLE001 - keep the poller alive
            state.last_error = str(exc)
            log.warning("HDX conflict stats fetch failed: %s", exc)
            await storage.record_source_health("hdx_conflict_stats", None, False, str(exc))
        # An empty exception message (e.g. a bare asyncio.TimeoutError) is
        # still a failure -- branch on whether the fetch itself succeeded,
        # not on the truthiness of the resulting error string.
        consecutive_failures = 0 if ok else consecutive_failures + 1
        await asyncio.sleep(POLL_INTERVAL if ok else min(FAILURE_RETRY_INTERVAL * consecutive_failures, POLL_INTERVAL))
